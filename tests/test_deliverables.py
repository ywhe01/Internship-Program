from pathlib import Path

from docx import Document
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]


def test_required_deliverables_exist():
    required = [
        ROOT / "README.md",
        ROOT / "demo" / "app.py",
        ROOT / "final" / "邻校掌柜_一页A4提交版.docx",
        ROOT / "assets" / "workflow.png",
        ROOT / "assets" / "chat-demo.png",
        ROOT / "assets" / "campaign-poster.png",
        ROOT / "assets" / "validation-plan.png",
        ROOT / "assets" / "dashboard.png",
    ]
    assert all(path.exists() and path.stat().st_size > 0 for path in required)


def test_document_is_a4_and_contains_required_sections():
    path = ROOT / "final" / "邻校掌柜_一页A4提交版.docx"
    document = Document(path)
    section = document.sections[0]
    assert abs(section.page_width.cm - 21.0) < 0.05
    assert abs(section.page_height.cm - 29.7) < 0.05

    texts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.extend(p.text for p in cell.paragraphs)
                for nested in cell.tables:
                    for nested_row in nested.rows:
                        for nested_cell in nested_row.cells:
                            texts.extend(p.text for p in nested_cell.paragraphs)
    joined = "\n".join(texts)
    for phrase in ("机会判断", "MVP", "商业判断", "两周验证"):
        assert phrase in joined


def test_visual_assets_have_expected_resolution():
    expected = {
        "workflow.png": (1600, 900),
        "chat-demo.png": (900, 1220),
        "campaign-poster.png": (900, 1200),
        "validation-plan.png": (1600, 900),
        "dashboard.png": (1600, 900),
    }
    for name, size in expected.items():
        with Image.open(ROOT / "assets" / name) as image:
            assert image.size == size
