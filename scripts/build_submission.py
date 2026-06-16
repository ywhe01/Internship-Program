from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "assets"
OUTPUT = ROOT / "final" / "邻校掌柜.docx"

NAVY = "18324A"
TEAL = "0E7C74"
ORANGE = "F28C45"
INK = "20262D"
MUTED = "5F6D75"
LINE = "D2DBDE"
WHITE = "FFFFFF"
PALE_ORANGE = "FFF1E7"


def shade(cell, fill: str):
    properties = cell._tc.get_or_add_tcPr()
    node = properties.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        properties.append(node)
    node.set(qn("w:fill"), fill)


def margins(cell, top=70, start=90, bottom=70, end=90):
    properties = cell._tc.get_or_add_tcPr()
    node = properties.first_child_found_in("w:tcMar")
    if node is None:
        node = OxmlElement("w:tcMar")
        properties.append(node)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        item = node.find(qn(f"w:{side}"))
        if item is None:
            item = OxmlElement(f"w:{side}")
            node.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")


def border(cell, color=LINE, size=6):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for side in ("top", "left", "bottom", "right"):
        item = borders.find(qn(f"w:{side}"))
        if item is None:
            item = OxmlElement(f"w:{side}")
            borders.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), str(size))
        item.set(qn("w:color"), color)


def clear(cell):
    cell.text = ""
    return cell.paragraphs[0]


def run(paragraph, text, size=8.2, bold=False, color=INK):
    item = paragraph.add_run(text)
    item.font.name = "Microsoft YaHei"
    item._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    item.font.size = Pt(size)
    item.bold = bold
    item.font.color.rgb = RGBColor.from_string(color)
    return item


def paragraph(paragraph, before=0, after=0, line=1.0, align=None):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def section_label(cell, text, color=TEAL):
    shade(cell, color)
    margins(cell, top=55, bottom=55)
    p = clear(cell)
    run(p, text, 8.5, True, WHITE)


def bullets(cell, items: list[str], size=7.8):
    margins(cell, top=60, bottom=55, start=95, end=85)
    for index, text in enumerate(items):
        p = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph(p, after=1.2, line=1.02)
        run(p, "• ", size, True, ORANGE)
        run(p, text, size, False, INK)


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(0.75)
    section.bottom_margin = Cm(0.65)
    section.left_margin = Cm(0.85)
    section.right_margin = Cm(0.85)
    section.header_distance = Cm(0.2)
    section.footer_distance = Cm(0.2)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(8)
    normal.paragraph_format.space_after = Pt(0)

    heading = document.add_paragraph()
    paragraph(heading, after=0.5)
    run(heading, "邻校掌柜", 17, True, NAVY)
    run(heading, "  |  校园周边餐饮店的 AI 增长值班经理", 12.5, True, TEAL)

    subheading = document.add_paragraph()
    paragraph(subheading, after=3)
    run(
        subheading,
        "老板发一句话，3分钟得到“今天做什么、是否赚钱、素材怎么发、核销后怎么复盘”。",
        8.5,
        False,
        MUTED,
    )

    evidence = document.add_table(rows=1, cols=3)
    evidence.autofit = False
    stats = [
        ("62万家", "2024年高校商圈社会餐饮门店[1]"),
        ("83.9%", "1140名大学生样本有外卖行为[2]"),
        ("98%", "餐饮企业属于中小微企业[3]"),
    ]
    for cell, (value, label) in zip(evidence.rows[0].cells, stats):
        shade(cell, NAVY)
        margins(cell, top=55, bottom=60)
        p = clear(cell)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(p, value, 13, True, WHITE)
        p2 = cell.add_paragraph()
        paragraph(p2, align=WD_ALIGN_PARAGRAPH.CENTER)
        run(p2, label, 6.7, False, "DBE6EA")

    spacer = document.add_paragraph()
    paragraph(spacer, after=0)

    main = document.add_table(rows=1, cols=2)
    main.autofit = False
    main.columns[0].width = Cm(12.0)
    main.columns[1].width = Cm(7.1)
    left, right = main.rows[0].cells
    margins(left, top=0, bottom=0, start=0, end=100)
    margins(right, top=0, bottom=0, start=100, end=0)

    label = left.add_table(rows=1, cols=1)
    section_label(label.cell(0, 0), "01 机会判断｜先服务一类能被核销验证的店")
    bullets(
        left,
        [
            "目标用户：高校500米内、1–3家店、老板亲自经营且没有专职运营的餐饮/茶饮店。",
            "痛点：客流随天气、课表、考试周和假期波动；促销靠经验，常见“有流水、没利润”；跨平台内容耗时且无法归因。",
            "用户信号：平台已观察到商家“想用AI却用不起、用不好”[4]；高校周边商家公开反馈营收下滑与价格战[5]。这些是二手证据，首周用15位老板访谈验证。",
            "现在值得做：多模态模型已能读取菜单/库存截图并低成本产出素材；本地生活具备内容→交易→核销闭环。",
        ],
    )

    label = left.add_table(rows=1, cols=1)
    section_label(label.cell(0, 0), "02 MVP｜不是文案工具，而是每日一条经营动作", NAVY)
    bullets(
        left,
        [
            "输入：老板语音 + 菜单/库存/昨日订单；补充天气、校历与校园节奏。",
            "决策：先算安全底价、单份贡献毛利与限量，再选择清库存/拉新/复购/提客单。",
            "执行：生成海报、群文案、15秒短视频分镜和专属核销码；发布前由老板确认。",
            "复盘：记录核销、自然销量基线和额外成本，计算增量贡献毛利并给出次日建议。",
        ],
    )

    label = left.add_table(rows=1, cols=1)
    section_label(label.cell(0, 0), "03 商业判断｜只为可归因的价值收费")
    bullets(
        left,
        [
            "付费方：门店老板；为节省运营时间和可归因的增量贡献毛利付费。",
            "种子验证69元/店；验证后199元/月含30次行动。首期不按GMV抽成，避免增量归因争议。",
            "获客：学生代理逐店演示，与POS服务商、打印店、餐饮供应商合作；主张“7天赚回月费”，不卖抽象AI能力。",
            "差异化：校园时空信号 + 门店毛利约束 + 跨平台素材 + 独立核销码，模型只负责表达，不直接决定价格。",
        ],
    )

    right_label = right.add_table(rows=1, cols=1)
    section_label(right_label.cell(0, 0), "一次完整行动", ORANGE)
    picture = right.add_paragraph()
    paragraph(picture, before=2, after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    picture.add_run().add_picture(str(ASSETS / "chat-demo.png"), width=Cm(5.65))

    plan_label = right.add_table(rows=1, cols=1)
    section_label(plan_label.cell(0, 0), "两周验证与停止条件", NAVY)
    bullets(
        right,
        [
            "D1–3：15位老板+30名学生访谈。",
            "D4–6：5店人工MVP，每店发布≥3次。",
            "D7–10：专属码A/B测试，记录核销与贡献毛利。",
            "D11–14：收取69元真实订金，复访拒付原因。",
            "GO：≥3/5店付费；素材采用率≥60%；日操作≤5分钟；7天增量毛利≥月费2倍。",
            "资源：1产品运营+1兼职开发+5名学生地推；两周预算约3000元。",
        ],
        7.25,
    )

    conclusion = document.add_table(rows=1, cols=1)
    cell = conclusion.cell(0, 0)
    shade(cell, PALE_ORANGE)
    margins(cell, top=50, bottom=50)
    p = clear(cell)
    run(p, "关键判断：", 8.1, True, ORANGE)
    run(
        p,
        "先验证老板是否愿为可归因的增量毛利付费。若核销无法证明增量，即使内容很好看，也不进入产品开发。",
        8.1,
        True,
        INK,
    )

    sources = document.add_paragraph()
    paragraph(sources, before=1, line=0.9)
    run(
        sources,
        "来源：[1] 红餐《中国高校餐饮研究报告2025》  [2] 马艳梅等《大学生外卖消费行为》  "
        "[3] 国家信息中心《中国餐饮业数字化发展报告2024》  [4] 美团商家AI公开信息  "
        "[5] 钛媒体高校周边商家走访。完整链接见 docs/evidence-and-sources.md。"
        "桌面研究日期：2026-06-15；未虚构一手访谈。",
        6.25,
        False,
        MUTED,
    )

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                border(cell, color=WHITE, size=0)

    document.save(OUTPUT)
    print(OUTPUT.relative_to(ROOT))


if __name__ == "__main__":
    build()
